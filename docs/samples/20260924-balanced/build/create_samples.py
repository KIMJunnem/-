from pathlib import Path
import sys, json, re, shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parent.parent
OUT=ROOT/'ready'; BUILD=ROOT/'build'
OUT.mkdir(parents=True,exist_ok=True)
FONTROOT=Path('C:/Users/vdfr7/Documents/Codex/Swan-Assets/fonts')
FONTS={'Pretendard':FONTROOT/'Pretendard/Pretendard-Regular.ttf','MaruBuri':FONTROOT/'MaruBuri/MaruBuri-Regular.ttf','IBM Plex Sans KR':FONTROOT/'ibmplexsanskr/IBMPlexSansKR-Regular.ttf'}
if '--fonts-only' in sys.argv:
 im=Image.new('RGB',(1500,780),'white');draw=ImageDraw.Draw(im)
 for i,(family,path) in enumerate(FONTS.items()):
  y=30+i*250
  draw.text((40,y),family,font=ImageFont.truetype(str(path),24),fill='#777777')
  draw.text((40,y+45),'흩어진 자료를 읽기 좋은 문서로',font=ImageFont.truetype(str(path),47),fill='black')
  draw.text((40,y+117),'회의에서 나온 결정과 미정 사항을 구분합니다. 21,000원 / 7장 / 32,000원',font=ImageFont.truetype(str(path),29),fill='black')
  draw.text((40,y+165),'무엇을 바꾸었고, 무엇을 남겼는지 한눈에 확인할 수 있습니다.',font=ImageFont.truetype(str(path),29),fill='black')
 im.save(BUILD/'font-comparison.jpg');sys.exit()

def document(kind):
 d=Document();s=d.sections[0]
 s.page_width=Inches(8.27);s.page_height=Inches(11.69)
 s.top_margin=Inches(.90);s.bottom_margin=Inches(.75);s.left_margin=s.right_margin=Inches(.82)
 s.header_distance=Inches(.32);s.footer_distance=Inches(.3)
 for g in list(s._sectPr.findall(qn('w:docGrid'))):s._sectPr.remove(g)
 for st in d.styles:
  rp=st.element.get_or_add_rPr(); rf=rp.find(qn('w:rFonts'))
  if rf is None:rf=OxmlElement('w:rFonts');rp.insert(0,rf)
  rf.attrib.clear()
  for key in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+key),'Pretendard')
  pp=st.element.find(qn('w:pPr'))
  if pp is not None:
   for b in list(pp.findall(qn('w:pBdr'))):pp.remove(b)
 for name,size in [('Normal',11),('Title',30),('Heading 1',22),('Heading 2',15)]:
  st=d.styles[name];st.font.name='Pretendard';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
  st.paragraph_format.line_spacing=1.35;st.paragraph_format.space_after=Pt(10)
  if name!='Normal':st.font.bold=True;st.paragraph_format.space_before=Pt(17)
 h=s.header.paragraphs[0];h.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 r=h.add_run('swan · SAMPLE');r.bold=True;r.font.size=Pt(18);r.font.color.rgb=RGBColor.from_string('B0B0B0')
 f=s.footer.paragraphs[0];r=f.add_run('swan  /  '+kind+'                                      ');r.font.size=Pt(8)
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
 d.core_properties.author='swan';d.core_properties.comments='자체 제작 서비스 시연 샘플'
 return d

def p(d,text,bold=False,size=11):
 q=d.add_paragraph();q.paragraph_format.line_spacing=1.4
 # Explicit word wrapping avoids broken Korean words in the preview renderer.
 font=ImageFont.truetype(str(FONTS['Pretendard']),round(size*8))
 lines=[]
 for para in text.split('\n'):
  line=''
  for word in para.split():
   candidate=(line+' '+word).strip()
   if line and font.getlength(candidate.replace('**',''))>6.63*72*8*.92:lines.append(line);line=word
   else:line=candidate
  lines.append(line)
 for i,part in enumerate(re.split(r'(\*\*.*?\*\*)','\n'.join(lines))):
  r=q.add_run(part[2:-2] if part.startswith('**') else part);r.bold=bold or part.startswith('**');r.font.size=Pt(size)
 return q
def title(d,kicker,text,desc):
 p(d,kicker,True,9);d.add_heading(text,0);p(d,desc,False,10)
def h(d,text):d.add_heading(text,2)
def table(d,heads,rows,widths):
 t=d.add_table(rows=1,cols=len(heads));t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for c,text in zip(t.rows[0].cells,heads):c.text=text
 for row in rows:
  for c,text in zip(t.add_row().cells,row):c.text=text
 for ri,row in enumerate(t.rows):
  for ci,c in enumerate(row.cells):
   c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr()
   borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
   pr.append(borders)
   pad=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    m=OxmlElement('w:'+side);m.set(qn('w:w'),'100');m.set(qn('w:type'),'dxa');pad.append(m)
   pr.append(pad)
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'222222' if ri==0 else ('F5F5F5' if ri%2==0 else 'FFFFFF'));pr.append(shade)
   for q in c.paragraphs:
    q.paragraph_format.space_after=Pt(3);q.paragraph_format.line_spacing=1.25
    for r in q.runs:r.font.size=Pt(10);r.bold=ri==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '111111')
 return t
def save(d,name):d.save(OUT/(name+'.docx'))

# 01 Provided-material writing: a realistic brand article, plus source brief.
d=document('일반 글 작성 샘플')
title(d,'WRITING  /  01','흩어진 자료를\n읽기 좋은 문서로','swan 서비스 소개 원고 · 제공 자료를 바탕으로 작성한 자체 샘플')
p(d,'메모는 충분한데 문서는 시작하기 어렵다면, 먼저 자료를 읽을 사람과 사용 목적을 정해 보세요. 같은 회의 메모도 팀 내부 공유용인지, 외부 제안용인지에 따라 첫 문장과 목차가 달라집니다.',size=12)
h(d,'읽는 사람이 알아야 할 내용부터')
p(d,'swan의 문서 작업은 보내주신 자료의 핵심을 찾는 데서 시작합니다. 결정된 사항과 아직 확인해야 할 내용을 나누고, 독자가 궁금해할 순서에 맞춰 문장을 이어 갑니다. 원자료에 없는 실적이나 수치는 확인 없이 넣지 않습니다.')
h(d,'글의 목적에 맞춘 구성')
p(d,'서비스 소개문에는 누구에게 어떤 도움을 주는지 먼저 씁니다. 내부 보고서에는 현재 상황과 판단 근거, 다음 할 일을 구분합니다. 이미 작성한 글이 있다면 내용을 새로 쓰기 전에 중복 문장과 문단 순서부터 살펴봅니다.')
h(d,'문서와 발표 자료를 함께 준비할 때')
p(d,'보고서를 PPT로 옮길 때는 문단을 그대로 붙여 넣지 않습니다. 한 장에서 설명할 내용을 정하고, 근거가 필요한 부분은 표나 수치로 정리합니다. 발표 자료를 본 사람이 원문으로 돌아가 세부 내용을 확인할 수 있도록 주요 용어를 맞춥니다.')
p(d,'**사용할 곳과 마감일, 가지고 계신 자료를 보내주세요.** 작업 범위를 확인한 뒤 필요한 분량과 전달 형식을 안내드립니다.',size=12)
d.add_page_break();title(d,'SOURCE BRIEF  /  02','제공 자료와 작성 기준','이 샘플의 원자료와 편집 판단을 함께 보여드립니다')
h(d,'작성에 사용한 자체 메모')
p(d,'브랜드 이름은 swan. 흑백 중심의 간결한 인상. 일반 글 작성, 문서 편집, 보고서의 PPT 변환, 한국어 자막 작업. 고객이 준 자료 중심. 자료에 없는 숫자는 확인. 전화보다 플랫폼 채팅. 샘플은 실제 고객 실적으로 소개하지 않기.')
table(d,['원자료','완성 원고의 처리'],[
 ['자료가 흩어져 있음','첫 문단에서 독자의 상황으로 설명'],['일반 글과 편집 모두 제공','새 원고 작성과 기존 글 정리를 구분'],['PPT 변환','문단 복사와 발표용 구성을 구분'],['실적 수치 없음','경력 연수와 매출 증가율을 새로 만들지 않음']],[2.1,4.53])
h(d,'납품 범위 예시')
p(d,'소개 원고와 소제목, 문단 구성을 함께 제공합니다. 이 샘플은 외부 자료 조사나 검색 순위 성과를 포함하지 않습니다. 사진 제작, 광고 운영, 플랫폼 게시 대행은 별도 범위입니다.')
h(d,'확인할 부분')
p(d,'브랜드명과 서비스 범위가 원자료와 같은지, 문단마다 주제가 하나인지, 독자에게 필요한 다음 행동이 분명한지 확인합니다. 최종 길이와 어조는 고객의 사용처에 맞춰 합의합니다.')
save(d,'01_Writing_Brand_Article')

# 02 Precise proofreading, no altered scope or ungrounded certainty.
d=document('교정 교열 샘플')
title(d,'PROOFREADING  / 01','문장의 의미를 지키는 교정','가상의 행사 안내문 · 원문과 교정본을 같은 조건으로 비교')
h(d,'교정 전')
p(d,'이번 워크숍은 10월 7일 수요일 오후 2시에 진행될 예정입니다. 참여하시는 분들께서는 시작 10분전 까지 도착해 주시면 됩니다. 준비물은 노트북과 필기도구를 지참해 주세요. 신청하신 모든분에게 결과물을 제공해 드릴 예정이므로 사전에 이메일 주소를 정확하게 기입 부탁드립니다. 실습시간은 약 60분 정도이며 상황에 따라 달라 질 수 있습니다.')
h(d,'교정 후')
p(d,'워크숍은 10월 7일 수요일 오후 2시에 진행할 예정입니다. **시작 10분 전까지** 도착해 주세요. 노트북과 필기도구를 준비해 주세요. 신청하신 **모든 분께** 결과물을 제공할 예정이니, 이메일 주소를 정확하게 적어 주세요. 실습 시간은 **약 60분**이며, 진행 상황에 따라 **달라질 수 있습니다**.')
h(d,'유지한 정보')
p(d,'일시, 도착 기준, 준비물, 자료 제공 대상과 실습 시간은 그대로 두었습니다. 원문의 “예정”과 “달라질 수 있음”을 남겨 확정된 약속처럼 바꾸지 않았습니다.')
p(d,'자체 제작한 가상의 안내문입니다. 실제 모집 행사나 고객 작업 실적이 아닙니다.',size=9)
d.add_page_break();title(d,'EDIT NOTES  / 02','수정한 부분과 이유','오탈자 수정과 문장 다듬기를 구분합니다')
table(d,['원문','교정','이유'],[
 ['10분전 까지','10분 전까지','띄어쓰기 정리'],['모든분','모든 분','관형어와 명사를 구분'],['준비물은 … 지참해 주세요','… 준비해 주세요','주어와 서술어의 연결 정리'],['약 60분 정도','약 60분','중복된 근사 표현 축약'],['달라 질 수','달라질 수','본용언의 활용을 붙여 씀']],[2.05,2.18,2.4])
h(d,'고객 확인이 필요한 정보')
p(d,'원문에 없는 행사 연도, 장소, 취소 규정은 추가하지 않습니다. 날짜와 요일의 일치 여부는 연도를 받은 뒤 확인합니다. 근거가 없는 문장을 단정적으로 강화하지 않습니다.')
h(d,'교정과 새 작성의 구분')
p(d,'교정은 원문의 의도와 사실을 유지하며 문장을 다듬는 작업입니다. 새로운 자료를 조사하거나 내용을 늘리는 요청은 새 작성 범위로 구분합니다. 고친 문장의 이유가 필요하면 이와 같은 수정표를 함께 정리할 수 있습니다.')
save(d,'02_Proofreading_Before_After')

# 03 Formatting: exact original wording retained, only hierarchy changes.
raw_sections=[('목적','협업 문서의 수정 요청을 한곳에 모아 최종본 혼선을 줄인다.'),('적용 대상','외부 전달용 제안서와 내부 공유용 보고서에 적용한다.'),('파일 이름','프로젝트명_문서명_v01_날짜 형식을 사용한다.'),('검토 절차','작성자가 초안을 공유한다. 검토자는 페이지와 문장을 지정해 의견을 남긴다. 작성자는 반영 결과를 기록한다.'),('확인 사항','숫자와 출처, 표 제목, 페이지 번호, 첨부 파일을 확인한다.'),('최종 전달','승인된 파일만 전달 폴더에 둔다. 수정 이력은 내부 폴더에 보관한다.')]
d=document('문서 양식 편집 샘플');title(d,'LAYOUT  /  01','협업 문서 관리 안내','편집 전 · 아래 원문의 내용은 다음 페이지에서도 그대로 유지됩니다')
p(d,' '.join(k+': '+v for k,v in raw_sections))
h(d,'이 페이지에서 비교할 점');p(d,'항목이 한 문단에 모여 있어 필요한 규칙을 찾기 어렵습니다. 다음 페이지는 제목, 행간, 문단 간격만 조정한 편집본입니다. 날짜나 담당자, 승인 조건을 임의로 새로 넣지 않았습니다.')
d.add_page_break();title(d,'LAYOUT  /  02','협업 문서 관리 안내','편집 후 · 원문 보존형 양식 정리')
for k,v in raw_sections:h(d,k);p(d,v)
save(d,'03_Formatting_Before_After')

# 04 A real licensed interview excerpt, timed transcript and clean reading copy.
SRC=Path('C:/Users/vdfr7/Documents/Codex/2026-09-13/37/swan-samples-20260923/multilingual-v2/output')
srt=(SRC/'Swan_ko_60s_v2_ko.srt').read_text(encoding='utf8')
blocks=[]
for b in re.split(r'\n\s*\n',srt.strip()):
 lines=b.splitlines();blocks.append((lines[1].split(' --> ')[0],''.join([]) if len(lines)<3 else ' '.join(lines[2:])))
d=document('녹취 타이핑 샘플');title(d,'TRANSCRIPTION  / 01','인터뷰 발췌 기록','한국어 인터뷰 편집본 60초 · 시간 표시형 전사 시연')
p(d,'원음의 이야기를 확인하기 위한 읽기용 기록입니다. 시각은 원본 전체 영상이 아니라 동봉한 60초 편집본 기준입니다. 말더듬과 일부 군말을 정리한 자막 기반 기록이며, 발화 그대로의 축어록이나 인증 녹취록은 아닙니다.',size=10)
for stamp,text in blocks:
 p(d,'**'+stamp[:8]+'**  '+text,size=10.5)
p(d,'원작 Wikitongues / Teddy Nee · CC BY-SA 4.0 · 발췌 및 문장 정리',size=8)
d.add_page_break();title(d,'READING COPY  / 02','읽기용 정리본','같은 발언을 주제별로 묶었습니다. 원문의 경험과 계획을 구분합니다')
h(d,'자기소개')
p(d,'화자는 대만을 무척 좋아한다고 말한다. 서울 남쪽의 과천에서 왔으며, 본인과 아버지, 어머니, 동생으로 가족을 소개한다.')
h(d,'대만에서의 경험')
p(d,'대만에 두 번 왔고 현재가 두 번째 방문이라고 말한다. 첫 방문은 교환학생으로 왔으며, 친구를 사귀고 대만 음식을 먹어 본 경험을 이야기한다.')
h(d,'현재 계획')
p(d,'두 번째 방문은 워킹홀리데이 목적이다. 곧 한국어를 가르칠 계획이며 다른 일도 찾으려고 노력하고 있다고 말한다. 취업이 확정됐거나 일을 이미 시작했다고 바꾸지 않았다.')
h(d,'출처와 사용 조건')
p(d,'WIKITONGUES — Hanbid speaking Korean\n제작: Wikitongues / Teddy Nee\n원본: commons.wikimedia.org/wiki/File:WIKITONGUES-_Hanbid_speaking_Korean.webm\n라이선스: CC BY-SA 4.0 · creativecommons.org/licenses/by-sa/4.0/\n변경: 인터뷰 발췌 편집본의 한국어 기록 및 요약\n이 파생 기록에도 CC BY-SA 4.0을 적용합니다.',size=9)
p(d,'실제 고객의 개인정보나 비공개 대화를 사용하지 않았습니다. 원 제작자가 swan을 추천한다는 의미가 아닙니다.',size=9)
save(d,'04_Transcription_Interview')

# 05 Scope-defined original script, separate from footage production.
d=document('대본 구성 샘플');title(d,'SCRIPT  / 01','문서 편집을 소개하는\n60초 영상 대본','swan 자체 브랜드 설명용 기획 샘플 · 촬영과 영상 제작은 포함하지 않습니다')
p(d,'문서를 보내기 전에 무엇을 확인하는지 보여주는 짧은 영상입니다. 실제 원고를 가상의 메모로 대체해 개인정보 없이 시연하도록 구성했습니다.',size=12)
table(d,['구간','화면','내레이션'],[
 ['00–08초','메모가 한 문단에 모인 원고','자료는 모았는데, 어디부터 정리해야 할지 막막한가요?'],
 ['08–20초','제목과 본문을 구분한 같은 원고','먼저 읽을 사람과 문서의 목적을 정합니다. 필요한 정보가 앞에 오도록 순서를 바꿉니다.'],
 ['20–33초','원문과 수정본의 같은 문장을 비교','표현은 다듬되 숫자와 조건은 지킵니다. 확인되지 않은 내용은 따로 표시합니다.'],
 ['33–46초','문서의 핵심을 나눈 PPT 두 장','발표 자료가 필요하다면, 한 장에 하나의 주제가 보이도록 원고를 다시 구성합니다.'],
 ['46–56초','마감일과 필요한 파일 형식 표시','가지고 계신 자료와 사용할 곳, 필요한 날짜를 채팅으로 알려주세요.'],
 ['56–60초','흑백 swan 워드마크','작업 범위를 확인하고 안내드리겠습니다.']],[.85,2.4,3.38])
p(d,'시간은 편집 가이드입니다. 실제 녹음 길이에 맞춰 장면 길이를 조정합니다.',size=9)
d.add_page_break();title(d,'PRODUCTION NOTES  / 02','대본의 제작 범위','의뢰용 원고로 바로 검토할 수 있도록 화면과 말을 나눕니다')
h(d,'제공하는 결과')
p(d,'60초 구성안, 구간별 화면 설명, 내레이션 원고를 제공합니다. 실제 촬영, 성우 녹음, 음악 라이선스, 완성 영상 제작은 이 대본 샘플의 납품 범위가 아닙니다.')
h(d,'촬영 자료와 표시 원칙')
p(d,'화면에는 자체 제작한 가상 문서만 사용합니다. 실제 고객 이름, 회사 정보, 비공개 자료를 노출하지 않습니다. 전후 비교는 같은 원문을 사용하고, 매출이나 처리 속도 향상을 실제 성과처럼 제시하지 않습니다.')
h(d,'기획 단계에서 받을 자료')
p(d,'소개할 서비스, 영상을 볼 대상, 공개할 채널과 원하는 길이를 확인합니다. 촬영 가능한 공간과 이미 가진 화면 자료가 있으면 그 안에서 장면을 구성합니다. 제공되지 않은 장면을 확보된 자료처럼 취급하지 않습니다.')
h(d,'검수 기준')
p(d,'원고의 표현과 판매 범위가 맞는지 확인합니다. 구간별 화면이 실제로 준비 가능한지, 읽는 속도에 비해 문장이 긴지, 마지막 안내가 시청자에게 필요한 행동을 설명하는지 검토합니다. 대본은 원고 검토 후 별도 견적 대상입니다.')
save(d,'05_Script_60s_Brand')

(BUILD/'source-formatting.json').write_text(json.dumps(raw_sections,ensure_ascii=False,indent=2),encoding='utf8')
(BUILD/'authoring.json').write_text(json.dumps({'newDocuments':5,'font':'Pretendard','fontDecision':'본문 숫자와 편집 비교 가독성 및 기존 브랜드 통일. MaruBuri와 IBM Plex Sans KR 비교 후 선택.','expectedPages':2,'sourceTranscript':str(SRC/'Swan_ko_60s_v2_ko.srt'),'externalPaidApiCalls':0},ensure_ascii=False,indent=2),encoding='utf8')
print('\n'.join(str(p) for p in OUT.glob('*.docx')))
