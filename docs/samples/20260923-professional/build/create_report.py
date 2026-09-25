from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import ImageFont

R=Path(__file__).resolve().parent.parent
O=R/'ready'; O.mkdir(parents=True,exist_ok=True)
d=Document(); s=d.sections[0]
for grid in list(s._sectPr.findall(qn('w:docGrid'))):s._sectPr.remove(grid)
s.page_width=Inches(8.5);s.page_height=Inches(11)
s.top_margin=Inches(.85);s.bottom_margin=Inches(.7)
s.left_margin=s.right_margin=Inches(.8)
s.header_distance=Inches(.3);s.footer_distance=Inches(.3)
for sty in d.styles:
 rpr=sty.element.find(qn('w:rPr'))
 if rpr is not None:
  fonts=rpr.find(qn('w:rFonts'))
  if fonts is not None:
   fonts.attrib.clear()
   for key in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+key),'Pretendard')
 ppr=sty.element.find(qn('w:pPr'))
 if ppr is not None:
  for b in list(ppr.findall(qn('w:pBdr'))):ppr.remove(b)
for name,size in [('Normal',10.5),('Title',29),('Heading 1',21),('Heading 2',13)]:
 st=d.styles[name];st.font.name='Pretendard';st.font.size=Pt(size);st.font.color.rgb=RGBColor(17,17,17)
 st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Pretendard')
 st.paragraph_format.line_spacing=Pt(size*1.35);st.paragraph_format.space_after=Pt(7)
 if name!='Normal':
  st.font.bold=True;st.paragraph_format.space_before=Pt(11);st.paragraph_format.keep_with_next=True
h=s.header.paragraphs[0];h.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=h.add_run('swan · SAMPLE');r.font.size=Pt(19);r.font.bold=True;r.font.color.rgb=RGBColor.from_string('B5B5B5')
f=s.footer.paragraphs[0];r=f.add_run('swan  /  운영 전략 보고서  ·  자체 제작 샘플                                        ');r.font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
fontroot=Path('C:/Users/vdfr7/Documents/Codex/Swan-Assets/fonts/Pretendard')
def p(text,bold=False,size=10.5):
 font=ImageFont.truetype(str(fontroot/('Pretendard-Bold.ttf' if bold else 'Pretendard-Regular.ttf')),round(size*8))
 lines=[]
 for para in text.split('\n'):
  line=''
  for word in para.split():
   cand=(line+' '+word).strip()
   if line and font.getlength(cand)>6.9*72*8*.92:lines.append(line);line=word
   else:line=cand
  lines.append(line)
 q=d.add_paragraph();r=q.add_run('\n'.join(lines));r.bold=bold;r.font.size=Pt(size)
 q.paragraph_format.line_spacing=Pt(size*1.4)
 return q
def title(k,t):p(k,True,9);d.add_heading(t,0)
def h2(t):d.add_heading(t,2)
def table(headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers));t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for cell,txt in zip(t.rows[0].cells,headers):cell.text=txt
 for row in rows:
  for cell,txt in zip(t.add_row().cells,row):cell.text=txt
 for ri,row in enumerate(t.rows):
  for ci,cell in enumerate(row.cells):
   cell.width=Inches(widths[ci])
   for q in cell.paragraphs:
    font=ImageFont.truetype(str(fontroot/('Pretendard-Bold.ttf' if ri==0 else 'Pretendard-Regular.ttf')),76)
    lines=[];line=''
    for word in q.text.split():
     candidate=(line+' '+word).strip()
     if line and font.getlength(candidate)>(widths[ci]-.17)*72*8*.90:lines.append(line);line=word
     else:line=candidate
    lines.append(line);q.text='\n'.join(lines)
    q.paragraph_format.space_after=Pt(5);q.paragraph_format.space_before=Pt(5);q.paragraph_format.line_spacing=Pt(13)
    for run in q.runs:run.font.size=Pt(9.5);run.bold=(ri==0)
   tcPr=cell._tc.get_or_add_tcPr();border=OxmlElement('w:tcBorders');bt=OxmlElement('w:bottom');bt.set(qn('w:val'),'single');bt.set(qn('w:sz'),'5');bt.set(qn('w:color'),'222222' if ri==0 else 'DDDDDD');border.append(bt);tcPr.append(border)
 return t

title('OPERATIONS BRIEF  /  2026.09','고객 문의에서 결제까지\n상담 전환 개선안')
p('디지털 제작 서비스의 운영 책임자를 위한\n측정 기준과 실행 설계',True,16)
p('대상: 일반 문서 편집 · 보고서의 PPT 변환 · 영상 자막 제작\n문서 성격: swan 자체 제작 전문 보고서 샘플 / 실제 성과 분석 아님',False,9)
h2('핵심 판단')
p('주문이 적다는 결과만으로 가격이나 전문성을 원인으로 단정할 수 없다. 견적이 고객에게 도달했는지, 질문에 답했는지, 작업 범위가 합의됐는지를 먼저 분리해 확인해야 한다. 이 보고서는 실제 고객 로그를 분석한 결과가 아니라, 해당 판단을 가능하게 만드는 운영 설계다.')
h2('우선 실행할 세 가지')
table(['우선순위','실행','완료 증거'],[
 ['01  응대 누락','실제 고객 질문과 시스템 알림을 분리한다. 미답변 문의를 먼저 처리한다.','고객 질문 시각·답변 시각·발송 확인 기록'],
 ['02  결과물 이해','요청과 가장 가까운 샘플 1개에 포함 범위와 납품 형식을 붙인다.','열리는 샘플 링크와 검수 이력'],
 ['03  결제 전 불확실성','총액·마감·수정 범위를 한 번에 확인하고 결제 경로를 안내한다.','고객의 범위 동의와 결제 확인 기록']],[1.1,3.3,2.5])
h2('실행 전제와 중단 조건')
p('샘플은 공개 권한이 있는 자체 제작물만 사용한다. 납기는 원자료·작업량·현재 제작 대기열을 확인한 후 약속한다. 잘못된 금액, 누락된 질문, 깨진 링크가 발견되면 해당 자동 안내를 멈추고 수정한다. 이 문서 자체가 운영 설정을 변경하지는 않는다.')

d.add_page_break();title('01  /  MEASUREMENT','전환율의 분모부터 맞춘다')
p('같은 견적 집단을 같은 기간 관찰해야 단계별 이탈을 비교할 수 있다.',True,14)
p('아래는 계산 방법을 설명하기 위한 가상 사례다. 실제 swan 실적이나 업계 평균이 아니다. 단계별 순서를 모두 거쳤고 읽음 기록도 빠짐없이 남았다고 가정한다.',False,9.5)
table(['단계','건수','직전 단계 대비','견적 전체 대비'],[
 ['견적 발송 확인','120','—','100.0%'],['읽음 관측','72','72 ÷ 120 = 60.0%','60.0%'],['실제 고객 문의','18','18 ÷ 72 = 25.0%','15.0%'],['작업 범위 동의','12','12 ÷ 18 = 66.7%','10.0%'],['결제 확인','6','6 ÷ 12 = 50.0%','5.0%']],[2.4,.7,2.1,1.7])
h2('이 수치로 말할 수 있는 것과 없는 것')
p('가상 사례에서는 읽음 이후 문의 단계의 직전 단계 전환율이 가장 낮다. 그러나 그 사실만으로 샘플 부족이나 높은 가격이 원인이라고 말할 수는 없다. 고객의 실제 질문, 응대 누락, 서비스 적합성을 함께 확인한 뒤 가설을 세운다.')
h2('집계 규칙')
p('한 고객·한 요청·첫 견적을 기본 단위로 삼고 중복 발송은 새 견적으로 세지 않는다. 결제는 클릭이나 결제 의향이 아니라 결제 확인 기록으로 집계한다. 내부 테스트와 실패한 발송은 제외한다. [1]')
p('실제 운영에서는 문의 없이 결제하거나 읽음 기록 없이 답장하는 경로도 있을 수 있다. 따라서 전체 견적 대비 결제율을 주지표로 두고, 단계별 순차 전환은 보조지표로 본다. 읽음 미관측을 미열람으로 단정하지 않는다.')
h2('관찰 기간과 비용')
p('실험에서는 견적별 발송 후 7일을 동일하게 관찰하는 안을 사용한다. 7일 미경과 건은 미성숙 표본으로 따로 표시하고 기간 이후 결제는 별도 추적한다. 7일은 본 샘플의 운영 제안이며 최적 기간이라는 근거는 없다.')
p('건별 기여이익 = 환불 반영 결제액 − 플랫폼·결제 수수료 − API·외주비 − 직접 투입 인건비. 고정비와 세금까지 반영한 순이익과 구분한다.',True,10)

d.add_page_break();title('02  /  RESPONSE DESIGN','고객이 묻는 것부터 답한다')
p('확인 질문을 늘리기 전에, 고객의 현재 결정에 필요한 답을 먼저 준다.',True,14)
table(['상황','응대 순서','전송 전 확인'],[
 ['언제 가능한가','현재 가능한 일정 → 일정 산정 조건 → 꼭 필요한 자료 요청','제작 대기열·파일 분량·난이도'],
 ['작업 경험이 있는가','확인된 경험만 설명 → 유사 자체 샘플 → 다른 점 명시','허위 경력·타 고객 자료 여부'],
 ['얼마인가','확인된 총액 → 포함·제외 범위 → 추가금 발생 조건','서비스·분량·옵션 일치'],
 ['진행하고 싶다','작업 조건 요약 → 동의 확인 → 플랫폼 결제 안내','총액·마감·수정·납품 형식']],[1.3,3.5,2.1])
h2('실제 응대 형태의 예시')
p('상황: 고객이 보고서를 발표 자료로 바꾸는 데 얼마나 걸리는지 문의. 원본은 아직 받지 않은 상태.',False,9.5)
p('“원본 분량을 아직 못 봐서 완료 시각은 파일을 확인한 뒤 말씀드릴게요. 보고서와 필요한 슬라이드 수, 발표 시간을 보내주시면 가능한 일정부터 먼저 확인하겠습니다. 아래는 같은 방식으로 정리한 자체 제작 샘플입니다.”',True,11)
p('위 문장은 미확인 상태의 예시다. 실제 가능 일정을 이미 확인했다면 그 일정을 첫 문장에 제시한다. 파일을 받았는데 다시 요청하거나, 확인하지 않은 당일 납품을 약속하지 않는다.',False,9.5)
h2('샘플이 설명해야 할 세 가지')
p('입력: 어떤 원자료를 받는가. 변환: 정보의 순서·표현·시각 구조를 어떻게 바꾸는가. 결과: 고객이 어떤 파일을 받아 어디에 쓸 수 있는가. 요청과 무관한 포트폴리오를 여러 개 보내지 않고 가장 가까운 한 사례부터 제시한다.')
h2('결제와 납품을 구분한다')
p('범위 합의 후 총액과 납기를 다시 확인하고 플랫폼의 정상 결제 경로를 안내한다. 결제 확인 전 최종 파일은 전달하지 않는다. 납품은 내용 대조·파일 열림·시각 검수를 통과한 버전에 한한다. 샘플에는 swan · SAMPLE을, 계약 납품본에는 합의된 표시만 적용한다.')

d.add_page_break();title('03  /  VALIDATION','한 번에 하나만 바꿔 검증한다')
p('검증 가설: 요청과 가까운 샘플을 견적과 함께 제공하면,\n결과물 이해가 높아져 결제 전환이 개선되는가?',True,14)
table(['항목','실험 설계안'],[
 ['배정','동일 서비스 안에서 고객 단위로 A/B 무작위 배정. 같은 고객은 같은 군 유지.'],
 ['변경 변수','A: 기존 견적 / B: 같은 견적 + 관련 자체 샘플 1개. 가격·범위·후속 연락 정책 동일.'],
 ['주지표','발송 확인된 견적 중 7일 이내 결제 확인 비율. 환불·기여이익은 별도 추적.'],
 ['보조·안전 지표','실제 답장률, 문의 후 응답 시간, 잘못된 안내, 수신 거부·불만.'],
 ['판단','먼저 누락 없이 기록 가능한지 소규모 점검. 기준 전환율과 최소 의미 차이를 정한 뒤 표본수·종료 기준을 사전 확정.']],[1.25,5.65])
h2('소수 주문으로 승자를 선언하지 않는다')
p('견적 10건에서 결제 2건이 나왔다면 관측 비율은 20%다. 이를 재현 가능한 성공률이나 샘플 제공의 효과로 단정하지 않는다. 실험 도중 수치가 좋아 보인다는 이유만으로 종료 시점을 바꾸지 않는다. 충분한 표본이 없으면 운영상 관찰로 남긴다.')
h2('책임과 실행 순서')
p('운영 담당: 미답변·중복·실패 발송을 매일 점검. 제작 담당: 샘플과 실제 제공 범위의 일치 확인. 검수 담당: 최종 렌더링과 링크·워터마크 확인. 운영 책임자: 주 1회 성숙 표본만 검토하고 다음 변경 1개를 결정. 소규모 조직은 겸임할 수 있으나 확인 기록은 구분한다.')
h2('근거와 적용 한계')
p('[1] GOV.UK Service Manual, Measuring completion rate (2021-02-19 갱신). 완료/시작의 비율, 테스트 제외, 시작·종료 정의 원칙을 참고했다. 정부 서비스의 기준을 민간 제작 상담에 맞게 응용했으며 정부의 전환율 목표를 인용한 것은 아니다.',False,8.5)
p('https://www.gov.uk/service-manual/measuring-success/measuring-completion-rate',False,8)
p('그 외의 기간·응대 절차·실험 설계는 본 샘플의 제안이다. 실제 고객 로그, 매출, 시장 평균을 검증한 보고서가 아니며, 적용 결과에 따라 수정해야 한다.',False,9)
d.core_properties.title='고객 문의에서 결제까지 — 상담 전환 개선안';d.core_properties.author='swan'
d.save(O/'Swan_Professional_Report.docx')
print(O/'Swan_Professional_Report.docx')
