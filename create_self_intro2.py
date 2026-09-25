from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pathlib import Path
out=Path('sales/테스트_자기소개서_릴레이데스크_2차.docx')
d=Document(); s=d.sections[0]; s.top_margin=Inches(.7); s.bottom_margin=Inches(.7); s.left_margin=Inches(.85); s.right_margin=Inches(.85)
for n,sz,bold,col in [('Normal',10.5,False,(32,35,43)),('Title',24,True,(32,35,43)),('Heading 1',15,True,(91,70,168))]:
 st=d.styles[n]; st.font.name='맑은 고딕'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'맑은 고딕'); st.font.size=Pt(sz); st.font.bold=bold; st.font.color.rgb=RGBColor(*col)
def ptext(t,style=None,after=7):
 p=d.add_paragraph(style=style); p.add_run(t); p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=1.35; return p
ptext('자기소개서 테스트 2차본','Title',4); ptext('지원 직무  업무 자동화 운영',None,16)
ps=[
('성장 과정과 가치관을 소개해 주세요.','저는 문제를 빠르게 덮기보다, 왜 그런 현상이 생겼는지 확인한 뒤 다시 같은 문제가 생겨도 대응할 수 있는 기준을 만드는 태도를 중요하게 생각합니다. 릴레이 데스크를 개선하는 과정에서도 답변 생성, 고객 채팅 전송, 파일 전달, 결제와 검수를 각각 나누어 보고, 한 단계의 성공을 다음 단계의 성공으로 착각하지 않으려 했습니다. 이런 태도는 단순히 꼼꼼하다는 말보다, 사용자가 실제로 겪는 불편을 기준으로 일하는 방식에 가깝습니다. 앞으로도 속도만 앞세우지 않고 목표와 조건을 먼저 정리해 신뢰할 수 있는 결과를 만들겠습니다.'),
('문제를 해결했던 경험과 그 과정에서 배운 점을 작성해 주세요.','자동화 흐름을 점검하면서 채팅이 특정 고객 화면만 반복해서 열리고, 결과 파일이 준비되어도 전송 여부를 확인하기 어려운 문제를 확인했습니다. 저는 이를 한 가지 오류로 처리하지 않고 고객 메시지 수집, 답변 생성, 화면 전송, 전송 확인의 단계로 나누어 살폈습니다. 특히 전송 결과가 불확실할 때 같은 파일을 반복 보내지 않도록 보류 상태를 두고, 실제 작업물과 검수 메모를 구분해야 한다는 기준을 세웠습니다. 이 과정에서 문제 해결의 핵심은 빠른 재시도보다 상태를 정확히 기록하고 실패 후의 다음 행동을 정하는 데 있다는 점을 배웠습니다.'),
('지원 직무에서 본인의 강점을 어떻게 발휘하겠습니까?','업무 자동화 운영 직무에서 저는 요구사항을 실행 가능한 단계와 확인 기준으로 바꾸는 강점을 발휘하겠습니다. 예를 들어 고객 요청은 목적·범위·파일 형식·기한으로 정리하고, 결과물은 내용뿐 아니라 파일이 실제로 열리는지, 이름과 버전이 맞는지, 전달이 확인됐는지까지 점검하겠습니다. 또한 문제가 생겼을 때 담당자에게 추측을 전달하기보다 확인된 사실과 추가 확인이 필요한 부분을 나누어 보고하겠습니다. 이를 통해 자동화의 처리 속도와 고객이 느끼는 신뢰를 함께 높이는 운영자가 되겠습니다.')]
for i,(q,a) in enumerate(ps,1): ptext(f'{i:02d}  {q}','Heading 1',8); p=ptext(a); p.paragraph_format.first_line_indent=Inches(.18)
ft=d.sections[0].footer.paragraphs[0]; ft.alignment=WD_ALIGN_PARAGRAPH.RIGHT; r=ft.add_run('메로나 문서사무소 · 내부 테스트 초안'); r.font.size=Pt(8); r.font.color.rgb=RGBColor(120,125,135)
d.core_properties.title='자기소개서 테스트 2차본'; d.core_properties.author='메로나 문서사무소'; d.save(out); print(out)
