from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from create_document_templates import setup, footer, heading, body, bullet, info_table, set_cell_text as cell_text, borders, shade
import json

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "artifacts" / "document-templates" / "rotation-pack"
OUT.mkdir(parents=True, exist_ok=True)

STYLES = {
    "editorial": {"title": 31, "h1": 18, "cover_align": WD_ALIGN_PARAGRAPH.LEFT, "table": "111111", "label": "EDITORIAL"},
    "executive": {"title": 27, "h1": 16, "cover_align": WD_ALIGN_PARAGRAPH.CENTER, "table": "2F2F2F", "label": "EXECUTIVE"},
    "minimal": {"title": 25, "h1": 15, "cover_align": WD_ALIGN_PARAGRAPH.RIGHT, "table": "555555", "label": "MINIMAL"},
}

SPECS = [
    ("현황분석보고서", "보고서", ["요약", "현황", "원인 분석", "개선안", "실행 일정"]),
    ("성과결과보고서", "보고서", ["결과 요약", "목표와 범위", "주요 성과", "미달 항목", "다음 계획"]),
    ("시장조사보고서", "보고서", ["조사 개요", "시장 현황", "고객 특성", "경쟁 환경", "시사점"]),
    ("프로젝트완료보고서", "보고서", ["완료 요약", "수행 범위", "산출물", "이슈와 대응", "인계 사항"]),
    ("정책검토보고서", "보고서", ["검토 결론", "검토 배경", "쟁점", "대안 비교", "권고안"]),
    ("회사소개서", "소개서", ["한눈에 보기", "회사 소개", "핵심 역량", "주요 실적", "문의"]),
    ("서비스소개서", "소개서", ["서비스 개요", "고객 문제", "제공 범위", "진행 절차", "기대 효과"]),
    ("브랜드소개서", "소개서", ["브랜드 이야기", "핵심 가치", "고객 경험", "제품과 서비스", "연락처"]),
    ("전문가프로필", "소개서", ["전문가 소개", "주요 분야", "작업 방식", "대표 경험", "의뢰 안내"]),
    ("자료요약브리프", "요약", ["결론부터", "핵심 내용", "주요 근거", "주의할 점", "다음 행동"]),
    ("회의내용요약", "요약", ["회의 결론", "논의 사항", "결정 사항", "담당 업무", "후속 일정"]),
    ("기사리서치요약", "요약", ["핵심 요약", "주요 사실", "관점 비교", "확인 필요", "활용 방향"]),
    ("도서자료요약", "요약", ["전체 개요", "핵심 주장", "주요 개념", "인상적인 내용", "적용 아이디어"]),
    ("정보형블로그원고", "원고", ["독자의 문제", "핵심 정보", "구체적 방법", "확인 사항", "마무리"]),
    ("후기형블로그원고", "원고", ["경험 배경", "선택 이유", "진행 과정", "느낀 점", "추천 대상"]),
    ("지역서비스원고", "원고", ["지역과 대상", "서비스 특징", "이용 과정", "선택 기준", "문의 안내"]),
    ("제품소개원고", "원고", ["제품 한눈에 보기", "주요 특징", "사용 방법", "추천 대상", "구매 전 확인"]),
    ("카드뉴스원고", "원고", ["표지 문구", "문제 제기", "핵심 정보", "정리", "행동 유도"]),
    ("업무제안서", "제안서", ["제안 요약", "현재 문제", "제안 내용", "추진 계획", "기대 효과"]),
    ("협업제안서", "제안서", ["협업 목적", "상대방 이점", "역할 분담", "일정", "협의 사항"]),
    ("프로젝트기획안", "제안서", ["기획 배경", "목표", "세부 구성", "일정과 역할", "성과 기준"]),
    ("행사운영제안서", "제안서", ["행사 개요", "운영 방향", "프로그램", "인력과 일정", "안전과 점검"]),
    ("업무매뉴얼", "가이드", ["문서 목적", "적용 범위", "업무 절차", "예외 처리", "점검표"]),
    ("서비스이용가이드", "가이드", ["이용 전 확인", "신청 방법", "진행 단계", "결과 확인", "자주 묻는 질문"]),
    ("신입업무가이드", "가이드", ["업무 개요", "첫날 준비", "반복 업무", "보고 방법", "확인 목록"]),
    ("고객응대가이드", "가이드", ["응대 원칙", "문의 접수", "상황별 답변", "분쟁 예방", "인계 기준"]),
    ("주간업무보고", "업무", ["이번 주 요약", "완료 업무", "진행 업무", "문제와 지원", "다음 주 계획"]),
    ("회의록", "업무", ["회의 정보", "안건", "주요 논의", "결정 사항", "후속 업무"]),
    ("업무인수인계서", "업무", ["인계 개요", "진행 중 업무", "반복 일정", "파일과 계정", "주의 사항"]),
    ("프로젝트브리프", "업무", ["프로젝트 정의", "목표와 대상", "범위", "일정과 담당", "완료 기준"]),
]

def cover(doc, title, category, style_name, serial):
    cfg = STYLES[style_name]
    p = doc.add_paragraph(); p.alignment = cfg["cover_align"]; p.paragraph_format.space_after = Pt(48 if style_name == "editorial" else 34)
    r = p.add_run("SWAN"); r.bold = True; r.font.name = "Arial"; r.font.size = Pt(11)
    p = doc.add_paragraph(); p.alignment = cfg["cover_align"]; p.paragraph_format.space_after = Pt(8)
    r = p.add_run(f'{cfg["label"]}  {category.upper()}  {serial:02d}'); r.bold = True; r.font.name = "Arial"; r.font.size = Pt(8); r.font.color.rgb = RGBColor(100,100,100)
    p = doc.add_paragraph(title, style="Title"); p.alignment = cfg["cover_align"]
    p = doc.add_paragraph("내용이 먼저 보이고 읽는 흐름이 자연스러운 실무 문서")
    p.alignment = cfg["cover_align"]; p.paragraph_format.space_after = Pt(34)
    t = doc.add_table(rows=4, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for i, (k,v) in enumerate((("문서 목적","[목적]"),("제출 대상","[독자 또는 기관]"),("작성일","[YYYY MM DD]"),("작성자","[이름 또는 팀명]"))):
        cell_text(t.cell(i,0),k,bold=True,color="666666",size=9); cell_text(t.cell(i,1),v,size=10)
        if style_name != "minimal": shade(t.cell(i,0),"F2F2F2")
    borders(t, "D9D9D9"); doc.add_page_break()

def build(spec, idx):
    title, category, sections = spec
    style_name = list(STYLES)[(idx-1) % 3]; cfg = STYLES[style_name]
    d = Document(); setup(d, f"{title} 템플릿"); footer(d)
    d.styles["Title"].font.size = Pt(cfg["title"]); d.styles["Heading 1"].font.size = Pt(cfg["h1"])
    cover(d, title, category, style_name, idx)
    heading(d, sections[0]); body(d, "[이 문서의 목적과 가장 중요한 결론을 먼저 적습니다. 독자가 무엇을 이해하거나 결정해야 하는지 분명하게 설명합니다.]")
    heading(d, "핵심 포인트", 2)
    for n in range(1,4): bullet(d, f"[핵심 내용 {n}]")
    for n, section in enumerate(sections[1:], 1):
        heading(d, f"{n} {section}")
        body(d, "[제공 자료를 바탕으로 사실과 해석을 구분해 작성합니다. 필요한 배경, 근거, 조건을 자연스러운 문단으로 정리합니다.]")
        if n in (1,3):
            info_table(d, ["구분", "내용", "확인"], [["항목 1","[구체적인 내용]","[상태]"],["항목 2","[구체적인 내용]","[상태]"],["항목 3","[구체적인 내용]","[상태]"]], [30,100,30])
    heading(d, "최종 확인")
    for text in ("[누락된 정보가 없는지 확인]", "[수치와 고유명사 확인]", "[독자가 취할 다음 행동 확인]"): bullet(d,text)
    path = OUT / f"{idx:02d}_Swan_{title}_{style_name}.docx"; d.save(path)
    return {"id": f"SWAN-DOC-{idx:02d}", "category": category, "title": title, "style": style_name, "file": path.name, "sections": sections}

items = [build(spec, i) for i, spec in enumerate(SPECS, 1)]
manifest = {
    "version": 2,
    "count": len(items),
    "selectionPolicy": {
        "mode": "category_then_least_recently_used",
        "avoidSameTemplateForSameCustomer": True,
        "avoidImmediateRepeat": True,
        "fallback": "least_recently_used_global"
    },
    "templates": items
}
(OUT / "rotation-manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"created={len(items)} out={OUT}")
