from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "artifacts" / "document-templates"
OUT.mkdir(parents=True, exist_ok=True)

BLACK = "111111"
GRAY = "666666"
LIGHT = "F2F2F2"
PALE = "F7F7F7"
LINE = "D9D9D9"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(table, color=LINE, size="6"):
    tbl_pr = table._tbl.tblPr
    node = tbl_pr.find(qn("w:tblBorders"))
    if node is None:
        node = OxmlElement("w:tblBorders")
        tbl_pr.append(node)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = node.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            node.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def margins(cell, top=130, start=150, bottom=130, end=150):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = tc_mar.find(qn(f"w:{name}"))
        if el is None:
            el = OxmlElement(f"w:{name}")
            tc_mar.append(el)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, bold=False, color=BLACK, size=9.5, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Malgun Gothic"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    margins(cell)


def setup(doc, title):
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.top_margin = Mm(20)
    sec.bottom_margin = Mm(18)
    sec.left_margin = Mm(22)
    sec.right_margin = Mm(22)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Malgun Gothic"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(7)
    for name, size in (("Title", 29), ("Heading 1", 18), ("Heading 2", 13), ("Heading 3", 11)):
        st = styles[name]
        st.font.name = "Malgun Gothic"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
        st.font.color.rgb = RGBColor.from_string(BLACK)
        st.font.size = Pt(size)
        st.font.bold = True
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.space_after = Pt(7)
    styles["Title"].paragraph_format.space_before = Pt(0)
    styles["Title"].paragraph_format.space_after = Pt(16)
    core = doc.core_properties
    core.title = title
    core.subject = "Swan 문서 작성 납품 템플릿"
    core.author = "Swan"


def footer(doc):
    for sec in doc.sections:
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run("SWAN  DOCUMENT STUDIO")
        r.font.name = "Arial"
        r.font.size = Pt(8)
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(GRAY)


def cover(doc, eyebrow, title, subtitle, fields):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(52)
    r = p.add_run("SWAN")
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(BLACK)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(9)
    r = p.add_run(eyebrow.upper())
    r.font.name = "Arial"
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(GRAY)
    doc.add_paragraph(title, style="Title")
    p = doc.add_paragraph(subtitle)
    p.paragraph_format.space_after = Pt(42)
    for run in p.runs:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor.from_string(GRAY)
    t = doc.add_table(rows=len(fields), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    t.columns[0].width = Mm(35)
    t.columns[1].width = Mm(115)
    borders(t)
    for i, (k, v) in enumerate(fields):
        set_cell_text(t.cell(i, 0), k, bold=True, color=GRAY, size=9)
        set_cell_text(t.cell(i, 1), v, size=10)
        shade(t.cell(i, 0), LIGHT)
    doc.add_page_break()


def heading(doc, text, level=1):
    return doc.add_paragraph(text, style=f"Heading {level}")


def body(doc, text="[내용을 입력하세요]"):
    return doc.add_paragraph(text)


def bullet(doc, text="[핵심 항목]"):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def info_table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if widths:
        for col, width in zip(t.columns, widths):
            col.width = Mm(width)
    for j, h in enumerate(headers):
        set_cell_text(t.cell(0, j), h, bold=True, color="FFFFFF", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(t.cell(0, j), BLACK)
    for i, row in enumerate(rows, start=1):
        cells = t.add_row().cells
        for j, val in enumerate(row):
            set_cell_text(cells[j], val, size=9.2, align=WD_ALIGN_PARAGRAPH.LEFT if j else WD_ALIGN_PARAGRAPH.CENTER)
            if i % 2 == 0:
                shade(cells[j], PALE)
    borders(t)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def save_report():
    d = Document(); setup(d, "일반 보고서 템플릿"); footer(d)
    cover(d, "REPORT", "일반 보고서", "핵심 결론이 먼저 보이는 실무형 보고서", [("작성 목적", "[보고 목적]"), ("제출 대상", "[기관 또는 담당자]"), ("작성일", "[YYYY MM DD]"), ("작성자", "[이름 또는 팀명]")])
    heading(d, "보고서 개요")
    body(d, "이 문서는 [주제]의 현황과 핵심 쟁점을 정리하고, 독자가 판단하거나 실행해야 할 사항을 제시합니다. 가장 중요한 결론은 [한 문장 결론]입니다.")
    heading(d, "핵심 요약", 2)
    for x in ("[핵심 결론]", "[근거 또는 주요 수치]", "[권고 사항 또는 다음 단계]"): bullet(d, x)
    heading(d, "1 현황과 배경")
    body(d, "[독자가 알아야 할 배경과 현재 상황을 2~3개 문단으로 설명하세요. 사실과 해석을 구분하고 필요한 범위를 명확히 적습니다.]")
    heading(d, "주요 현황", 2)
    info_table(d, ["구분", "현재 상태", "의미"], [["항목 1", "[현황]", "[해석]"], ["항목 2", "[현황]", "[해석]"], ["항목 3", "[현황]", "[해석]"]], [30, 55, 75])
    heading(d, "2 분석")
    body(d, "[자료에서 확인한 패턴과 원인을 설명하세요. 단순 나열보다 항목 간 관계가 드러나도록 작성합니다.]")
    heading(d, "분석 기준", 2)
    info_table(d, ["기준", "확인 내용", "판단"], [["효과", "[내용]", "[판단]"], ["비용", "[내용]", "[판단]"], ["실행 가능성", "[내용]", "[판단]"]], [30, 90, 40])
    heading(d, "3 제안과 실행 계획")
    body(d, "[앞선 분석에 근거한 제안을 우선순위 순으로 제시합니다.]")
    info_table(d, ["순서", "실행 항목", "담당", "일정"], [["1", "[즉시 실행]", "[담당]", "[일정]"], ["2", "[단기 실행]", "[담당]", "[일정]"], ["3", "[후속 점검]", "[담당]", "[일정]"]], [18, 82, 30, 30])
    heading(d, "4 결론")
    body(d, "[핵심 결론을 다시 한 번 정리하고 독자가 취해야 할 다음 행동을 분명히 적습니다.]")
    d.save(OUT / "Swan_일반_보고서_템플릿.docx")


def save_intro():
    d = Document(); setup(d, "회사 서비스 소개서 템플릿"); footer(d)
    cover(d, "PROFILE", "회사 및 서비스 소개서", "짧고 설득력 있게 정리하는 소개 문서", [("회사 또는 브랜드", "[브랜드명]"), ("핵심 서비스", "[서비스명]"), ("문의", "[연락 방법]"), ("작성일", "[YYYY MM DD]")])
    heading(d, "한눈에 보는 소개")
    body(d, "[브랜드명]은 [고객 대상]에게 [핵심 서비스]를 제공합니다. 고객이 얻는 가장 분명한 변화는 [핵심 효익]입니다.")
    info_table(d, ["핵심 가치", "제공 내용", "고객 효익"], [["[가치 1]", "[서비스 설명]", "[효익]"], ["[가치 2]", "[서비스 설명]", "[효익]"], ["[가치 3]", "[서비스 설명]", "[효익]"]], [35, 65, 60])
    heading(d, "주요 서비스")
    for n in range(1, 4):
        heading(d, f"서비스 {n}  [서비스명]", 2)
        body(d, "[누구에게 필요한지, 무엇을 제공하는지, 결과물이 무엇인지 3~4문장으로 설명하세요.]")
    heading(d, "진행 절차")
    info_table(d, ["단계", "진행 내용", "고객 준비 사항"], [["01 상담", "범위와 일정을 확인합니다", "목적과 자료"], ["02 제작", "합의한 범위로 작성합니다", "추가 자료 회신"], ["03 검수", "완성도를 점검하고 수정합니다", "구체적인 피드백"], ["04 전달", "최종 형식으로 납품합니다", "최종 확인"]], [24, 76, 60])
    heading(d, "작업 기준")
    for x in ("제공 자료를 기준으로 사실관계를 정리합니다.", "작업 범위와 수정 횟수는 시작 전에 합의합니다.", "최종 결과물은 문장과 형식을 검수한 뒤 전달합니다."): bullet(d, x)
    heading(d, "문의")
    body(d, "[연락 방법]으로 목적, 분량, 보유 자료, 마감일을 알려주시면 작업 범위와 일정을 안내합니다.")
    d.save(OUT / "Swan_회사_서비스_소개서_템플릿.docx")


def save_brief():
    d = Document(); setup(d, "요약 브리프 템플릿"); footer(d)
    cover(d, "BRIEF", "핵심 요약 브리프", "긴 자료를 빠르게 이해하는 2페이지 요약", [("원문", "[자료명]"), ("독자", "[읽는 사람]"), ("요약 목적", "[의사결정 또는 공유]"), ("작성일", "[YYYY MM DD]")])
    heading(d, "결론부터")
    body(d, "[이 자료에서 독자가 가장 먼저 알아야 할 결론을 2~3문장으로 적습니다.]")
    heading(d, "핵심 내용", 2)
    for x in ("[핵심 내용 1]", "[핵심 내용 2]", "[핵심 내용 3]", "[핵심 내용 4]"): bullet(d, x)
    heading(d, "주요 근거")
    info_table(d, ["근거", "내용", "출처 위치"], [["1", "[주요 수치 또는 주장]", "[쪽 또는 항목]"], ["2", "[주요 수치 또는 주장]", "[쪽 또는 항목]"], ["3", "[주요 수치 또는 주장]", "[쪽 또는 항목]"]], [20, 105, 35])
    heading(d, "확인하거나 결정할 사항")
    for x in ("[결정할 사항]", "[추가 확인이 필요한 사항]", "[다음 행동]"): bullet(d, x)
    d.save(OUT / "Swan_핵심_요약_브리프_템플릿.docx")


def save_article():
    d = Document(); setup(d, "블로그 일반 원고 템플릿"); footer(d)
    cover(d, "ARTICLE", "블로그 및 일반 원고", "독자가 자연스럽게 끝까지 읽는 정보형 원고", [("주제", "[원고 주제]"), ("독자", "[주요 독자]"), ("목표", "[정보 제공 또는 행동 유도]"), ("분량", "[예상 글자 수]")])
    heading(d, "제목 후보")
    for x in ("[검색 의도를 반영한 제목]", "[독자의 문제를 드러내는 제목]", "[구체적인 결과를 보여주는 제목]"): bullet(d, x)
    heading(d, "도입")
    body(d, "[독자가 겪는 상황을 구체적으로 제시하고, 이 글을 읽으면 무엇을 알 수 있는지 안내합니다. 과장된 표현은 피하고 바로 본론으로 연결합니다.]")
    heading(d, "1 첫 번째 핵심 내용")
    body(d, "[핵심 정보를 설명합니다. 짧은 문단과 구체적인 예시를 사용합니다.]")
    heading(d, "2 두 번째 핵심 내용")
    body(d, "[앞 내용과 어떻게 이어지는지 설명하고 독자가 적용할 방법을 제시합니다.]")
    heading(d, "3 확인할 점")
    info_table(d, ["확인 항목", "설명"], [["[항목 1]", "[확인 방법]"], ["[항목 2]", "[확인 방법]"], ["[항목 3]", "[확인 방법]"]], [45, 115])
    heading(d, "마무리")
    body(d, "[핵심 내용을 한 번 정리하고 독자가 취할 수 있는 다음 행동을 자연스럽게 제안합니다.]")
    d.save(OUT / "Swan_블로그_일반_원고_템플릿.docx")


if __name__ == "__main__":
    save_report(); save_intro(); save_brief(); save_article()
    print(OUT)
