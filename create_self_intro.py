from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

out=Path('sales/테스트_자기소개서_릴레이데스크.docx')
doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(0.7); sec.bottom_margin=Inches(0.7); sec.left_margin=Inches(0.85); sec.right_margin=Inches(0.85)
styles=doc.styles
styles['Normal'].font.name='맑은 고딕'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'맑은 고딕'); styles['Normal'].font.size=Pt(10.5); styles['Normal'].font.color.rgb=RGBColor(32,35,43)
for name,size,color in [('Title',24,(32,35,43)),('Heading 1',15,(91,70,168)),('Heading 2',11.5,(32,35,43))]:
    st=styles[name]; st.font.name='맑은 고딕'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'맑은 고딕'); st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor(*color)

def shade(p, fill):
    pr=p._p.get_or_add_pPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); pr.append(shd)
def add_para(text='', style=None, bold_prefix=None):
    p=doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold=True; p.add_run(text[len(bold_prefix):])
    else: p.add_run(text)
    p.paragraph_format.space_after=Pt(7); p.paragraph_format.line_spacing=1.35
    return p
p=doc.add_paragraph(style='Title'); p.add_run('자기소개서 테스트 초안'); p.paragraph_format.space_after=Pt(4)
p=doc.add_paragraph(); r=p.add_run('지원자 특징 기반 · 기본형 3문항'); r.bold=True; r.font.size=Pt(10.5); r.font.color.rgb=RGBColor(91,70,168); p.paragraph_format.space_after=Pt(18)
intro=add_para('문제를 끝까지 확인하고, 고객이 실제로 사용할 수 있는 결과를 만들기 위해 기준을 세우고 개선해 온 태도를 중심으로 작성했습니다.')
shade(intro,'F5F2FB')
questions=[
('성장 과정과 가치관을 소개해 주세요.', '저는 한 번 정한 답을 그대로 받아들이기보다, 실제로 작동하는지 확인하고 더 나은 방법을 찾는 태도를 중요하게 생각합니다. 일을 진행할 때도 먼저 목표와 조건을 분명히 정리한 뒤, 결과를 직접 확인하며 부족한 부분을 고쳐 왔습니다. 특히 작은 오류라도 고객의 사용 경험을 바꿀 수 있다는 점을 의식해, “대략 이 정도면 됐다”에서 멈추지 않고 원인과 재발 가능성까지 살피려 합니다. 이 과정에서 제게 가장 중요한 기준은 속도와 완성도를 함께 지키는 것입니다. 빠르게 시작하되, 확인할수록 좋아지는 구조를 만드는 사람이 되고 싶습니다.'),
('문제를 해결했던 경험과 그 과정에서 배운 점을 작성해 주세요.', '자동화 시스템을 점검하는 과정에서 답변이 멈추거나 같은 화면만 반복되는 문제를 발견한 적이 있습니다. 처음에는 단순한 오류로 보였지만, 상태 확인과 실제 전송 결과가 분리되어 있고 실패 이후의 복귀 흐름이 충분히 기록되지 않는 것이 원인일 수 있다고 판단했습니다. 그래서 문제를 증상으로만 처리하지 않고, 수집·판단·전송·결과 확인의 단계를 나누어 확인하는 방식으로 접근했습니다. 각 단계에서 무엇이 성공했는지 증거를 남기고, 불확실한 경우에는 중복 실행을 막도록 기준을 세웠습니다. 이 경험을 통해 문제 해결은 한 번의 수정으로 끝나는 일이 아니라, 같은 문제가 다시 생겨도 원인을 찾을 수 있게 만드는 과정이라는 점을 배웠습니다.'),
('지원 직무에서 본인의 강점을 어떻게 발휘하겠습니까?', '저의 강점은 요구사항을 구체적인 실행 기준으로 바꾸고, 결과를 사용자의 관점에서 다시 확인하는 데 있습니다. 작업을 시작할 때 고객이 원하는 파일 형식과 범위를 먼저 정리하고, 제작 후에는 내용뿐 아니라 열림 여부, 파일명, 전달 상태까지 점검하는 방식을 선호합니다. 또한 한 번 만든 결과물도 피드백을 근거로 다시 개선할 수 있습니다. 지원한 직무에서도 업무를 작은 단계로 나누고, 확인 가능한 기록을 남기며, 필요한 경우 동료와 기준을 맞추겠습니다. 최종적으로는 맡은 일을 끝내는 데 그치지 않고, 다음 사람이 이어서 사용해도 혼선이 없는 결과를 만드는 구성원이 되겠습니다.')
]
for i,(q,a) in enumerate(questions,1):
    doc.add_paragraph(f'{i:02d}  {q}', style='Heading 1')
    p=add_para(a)
    p.paragraph_format.first_line_indent=Inches(0.18)
# footer
for section in doc.sections:
    footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    rr=footer.add_run('메로나 문서사무소 · 테스트 결과물'); rr.font.size=Pt(8); rr.font.color.rgb=RGBColor(120,125,135)
doc.core_properties.title='자기소개서 테스트 초안'; doc.core_properties.subject='기본형 3문항'; doc.core_properties.author='메로나 문서사무소'
doc.save(out)
print(out.resolve())
